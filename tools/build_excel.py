"""Erzeugt docs/Balancing.xlsx aus config/params.json und den Simulationsergebnissen.

Aufruf (im Repository-Ordner):
    node tools/simulate.mjs 60 --json docs/data/simulation.json
    node tools/sensitivity.mjs 30 --json docs/data/sensitivity.json
    python3 tools/build_excel.py
Geänderte Werte im Blatt «Parameter» übernimmt tools/excel_to_params.py zurück in config/params.json.
"""
import json
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
CFG = json.loads((ROOT / "config/params.json").read_text(encoding="utf-8"))
SIM = json.loads((ROOT / "docs/data/simulation.json").read_text(encoding="utf-8"))
SENS = json.loads((ROOT / "docs/data/sensitivity.json").read_text(encoding="utf-8"))

FONT = "Arial"
F = Font(name=FONT, size=10)
FB = Font(name=FONT, size=10, bold=True)
FH = Font(name=FONT, size=14, bold=True)
FINPUT = Font(name=FONT, size=10, color="0000FF")
FLINK = Font(name=FONT, size=10, color="008000")
FILL_IN = PatternFill("solid", fgColor="FFF2CC")
FILL_HEAD = PatternFill("solid", fgColor="DCEBEA")
THIN = Side(style="thin", color="D9D3C7")
BORDER = Border(bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

DESC = {
    "version": ("Allgemein", "Regelversion; bei Änderungen erhöhen (z. B. 2.2.1), damit Statistiken nur gleiche Regeln vergleichen", "Text"),
    "rounds": ("Allgemein", "Standard-Rundenzahl pro Spiel (wählbar 4–6)", "Runden"),
    "actionsPerRound": ("Allgemein", "Aktionen pro Team und Runde", "Anzahl"),
    "maxTeams": ("Allgemein", "Maximale Anzahl Teams", "Anzahl"),
    "upstreamCostPerParcel": ("Wirtschaftlichkeit", "Vorlaufkosten pro Paket (Abholung, Sortierung, Hauptlauf); nur für den Lieferketten-Anteil", "CHF/Paket"),
    "tariffPerParcel": ("Wirtschaftlichkeit", "Last-Mile-Tarif: Anteil des Paketpreises nach Abzug der Vorlaufkosten; Mindeststandard Kosten ≤ Tarif", "CHF/Paket"),
    "standards.onTimeNational": ("Mindeststandards", "Pünktlichkeit landesweit, mindestens", "Anteil"),
    "standards.onTimeZone": ("Mindeststandards", "Pünktlichkeit in jeder Zone, mindestens", "Anteil"),
    "timeWindowFailFactor": ("Zonen", "Faktor auf die Fehlzustellungsrate bei Zeitfenstern (0.5 = halbiert)", "Faktor"),
    "labour.baseWageGrowth": ("Arbeitsmarkt", "Lohnanstieg pro Runde ohne Knappheit", "Anteil"),
    "labour.scarcitySensitivity": ("Arbeitsmarkt", "Zusätzlicher Lohnanstieg pro 100 % mehr Fahrer als zu Spielbeginn (alle Teams)", "Faktor"),
}
ZONE_DESC = {
    "demand": ("Pakete pro Team, Feld und Runde", "Pakete"),
    "distance": ("Distanzfaktor auf variable Kosten und CO₂ von Fahrzeugen", "Faktor"),
    "failRate": ("Fehlzustellungsrate bei Haustürzustellung (Standardpolitik)", "Anteil"),
    "timeWindowCapacityLoss": ("Kapazitätsverlust eigener Fahrzeuge bei Zeitfenstern", "Anteil"),
}
EL_NAMES = {"vanDiesel": "Diesel-Van", "vanE": "E-Van", "hub": "Micro-Hub mit Cargo-Bikes", "box": "Paketbox/PUDO", "post": "Post-Vertrag", "crowd": "Crowd-Fahrer", "emergency": "Notfall-Subunternehmer"}
EL_DESC = {
    "driverCost": ("Lohn pro Fahrer und Runde (wird mit dem Lohnindex multipliziert)", "CHF/Runde"),
    "leasing": ("Leasing bzw. Miete pro Runde", "CHF/Runde"),
    "drivers": ("Benötigte Fahrer (zählen am gemeinsamen Arbeitsmarkt)", "Anzahl"),
    "varCost": ("Variable Kosten pro Paket", "CHF/Paket"),
    "co2": ("CO₂ pro Paket (bei Fahrzeugen mal Distanzfaktor)", "g/Paket"),
    "capacity": ("Kapazität pro Runde (in Stadt-Paketen)", "Einheiten"),
    "range": ("Reichweite ab Standortfeld", "Felder"),
    "resupplyCost": ("Nachschubkosten pro Paket (Zubringer zu Hub/Box)", "CHF/Paket"),
    "resupplyCo2": ("Nachschub-CO₂ pro Paket", "g/Paket"),
    "buildDelay": ("Bauzeit bis zur Wirkung", "Runden"),
    "acceptance": ("Anteil der Kundschaft, der an der Box abholt", "Anteil"),
    "failRate": ("Fehlzustellungsrate", "Anteil"),
    "feePerRegion": ("Vertragsgebühr pro Landregion und Runde (3 Regionen)", "CHF/Runde"),
    "onTimeRate": ("Anteil am nächsten Tag erfolgreich zugestellt (fiktiver Spielwert)", "Anteil"),
    "capacityShare": ("Kapazität als Anteil der normalen Regionsnachfrage", "Anteil"),
}
EV_NAMES = {"zez": "Null-Emissions-Zone", "peak": "Black Friday", "co2tax": "CO₂-Abgabe", "boxBoom": "Boxen beliebter", "strike": "Streik", "fuel": "Treibstoffpreis", "shortage": "Fahrermangel", "demandSurge": "Nachfrageschub"}
EV_DESC = {
    "announceLead": ("Vorlauf der Ankündigung", "Runden"),
    "demandFactor": ("Faktor auf die Nachfrage", "Faktor"),
    "chfPerKg": ("Abgabe pro kg CO₂ (fiktive Regulierung)", "CHF/kg"),
    "acceptance": ("Neue Akzeptanzquote der Box", "Anteil"),
    "capacityFactor": ("Faktor auf die Kapazität", "Faktor"),
    "duration": ("Dauer", "Runden"),
    "dieselVarFactor": ("Faktor auf die variablen Diesel-Kosten", "Faktor"),
    "emergencyExtra": ("Zuschlag Notfall-Subunternehmer", "CHF/Paket"),
    "wageJump": ("Einmaliger Lohnsprung", "Anteil"),
    "factor": ("Faktor auf die Nachfrage der betroffenen Zone", "Faktor"),
}


def leaves(obj, prefix=""):
    for k, v in obj.items():
        path = f"{prefix}{k}"
        if isinstance(v, dict):
            yield from leaves(v, path + ".")
        elif isinstance(v, (int, float, str)) and not isinstance(v, bool):
            yield path, v


def describe(path):
    if path in DESC:
        return DESC[path]
    parts = path.split(".")
    if parts[0] == "zones":
        d, u = ZONE_DESC.get(parts[2], (parts[2], ""))
        return (f"Zone {parts[1].capitalize()}", d, u)
    if parts[0] == "elements":
        name = EL_NAMES.get(parts[1], parts[1])
        if parts[2] == "weight":
            return (name, f"Kapazitätsverbrauch pro Paket in Zone {parts[3].capitalize()} (längere Wege = mehr)", "Faktor")
        d, u = EL_DESC.get(parts[2], (parts[2], ""))
        return (name, d, u)
    if parts[0] == "events":
        if len(parts) == 2:
            return ("Ereignisse", {"shocksPerGame": "Überraschungen pro Spiel", "extraAnnouncements": "Zusätzliche Ankündigungen (CO₂-Abgabe oder Boxen beliebter)"}.get(parts[1], parts[1]), "Anzahl")
        d, u = EV_DESC.get(parts[2], (parts[2], ""))
        return (f"Ereignis {EV_NAMES.get(parts[1], parts[1])}", d, u)
    return ("Sonstiges", path, "")


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = FB
        cell.fill = FILL_HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


wb = Workbook()

# ---------- Anleitung ----------
ws = wb.active
ws.title = "Anleitung"
lines = [
    ("Balancing «Redesign the Last Mile»", FH),
    (f"Regelversion {CFG['version']} · erstellt am {date(2026, 10, 29).strftime('%d.%m.%Y')} · Simulation mit {SIM['nSeeds']} Ereignisfolgen pro Strategie", F),
    ("", F),
    ("So arbeitet ihr mit dieser Mappe", FB),
    ("1. Blatt «Parameter»: nur die blauen Zahlen in der Spalte «Wert» (gelb hinterlegt) ändern. Jede Zahl hat eine Beschreibung und Einheit.", F),
    ("2. Blatt «Startnetz»: rechnet mit Formeln, was das unveränderte Diesel-Netz in einer normalen Runde kostet. Es reagiert sofort auf eure Parameter.", F),
    ("3. Blätter «Teststrategien», «Kalibrierung» und «Sensitivität»: Ergebnisse der automatischen Simulation mit dem echten Rechenkern. Sie ändern sich erst, wenn die Simulation neu läuft.", F),
    ("4. Geänderte Mappe an Claude geben (oder: python3 tools/excel_to_params.py docs/Balancing.xlsx). Dann wird config/params.json aktualisiert, die Simulation neu gerechnet und das Spiel neu veröffentlicht.", F),
    ("5. Bei jeder Änderung die Regelversion erhöhen (Parameter «version»), damit die Statistiken nur Spiele mit gleichen Regeln vergleichen.", F),
    ("", F),
    ("Farbcode", FB),
    ("Blaue Schrift auf Gelb = Eingabe · schwarze Schrift = Formel · grüne Schrift = Bezug auf ein anderes Blatt · graue Werte = Ergebnis der Simulation (kein Formelbezug)", F),
    ("", F),
    ("Herkunft der Zahlen", FB),
    ("Alle Werte sind bewusst vereinfachte Spielannahmen, kalibriert per Simulation (Kriterien im Blatt «Kalibrierung»). Sie sind keine Kennzahlen realer Unternehmen. Die WEF-Berichte (2020, 2024) und weitere Literatur sind für die finale Kalibrierung noch auszuwerten.", F),
]
for i, (txt, font) in enumerate(lines, start=1):
    c = ws.cell(row=i, column=1, value=txt)
    c.font = font
    c.alignment = WRAP
set_widths(ws, [130])

# ---------- Parameter ----------
wp = wb.create_sheet("Parameter")
wp.append(["Schlüssel", "Bereich", "Beschreibung", "Wert", "Einheit", "Quelle / Annahme"])
style_header(wp, 1, 6)
param_row = {}
for path, val in leaves(CFG):
    area, desc, unit = describe(path)
    wp.append([path, area, desc, val, unit, "Spielannahme, kalibriert per Simulation"])
    r = wp.max_row
    param_row[path] = r
    for c in range(1, 7):
        wp.cell(row=r, column=c).font = F
        wp.cell(row=r, column=c).alignment = WRAP
        wp.cell(row=r, column=c).border = BORDER
    vc = wp.cell(row=r, column=4)
    vc.font = FINPUT
    vc.fill = FILL_IN
    if unit == "Anteil":
        vc.number_format = "0.0%"
    elif unit.startswith("CHF"):
        vc.number_format = "0.00"
for p in ["elements.post.onTimeRate", "events.co2tax.chfPerKg"]:
    if p in param_row:
        wp.cell(row=param_row[p], column=6, value="Fiktiver Spielwert, keine Kennzahl eines realen Anbieters bzw. keine geltende Abgabe (CH: CO₂-Abgabe nur auf Brennstoffe)")
wp.cell(row=param_row["tariffPerParcel"], column=4).comment = Comment("Kalibriert: gute Strategien erfüllen alle Standards in rund 70 % der Ereignisfolgen; das Startnetz nie.", "Claude")
wp.freeze_panes = "A2"
set_widths(wp, [34, 26, 60, 12, 12, 40])
wp.cell(row=1, column=8, value="Startnetz: 4 Diesel-Vans in S1, je 2 in A2 und A8, je 1 im Zentrum jeder Landregion (in params.json unter startNetwork)").font = Font(name=FONT, size=9, italic=True)

P = lambda path: f"Parameter!$D${param_row[path]}"  # noqa: E731

# ---------- Startnetz (Formeln) ----------
wsn = wb.create_sheet("Startnetz")
wsn["A1"] = "Kostenstruktur des Startnetzes in einer normalen Runde (ohne Ereignisse)"
wsn["A1"].font = FH
wsn["A2"] = "Vereinfachte Rechnung mit Formeln: Fixkosten nach Standortzone, variable Kosten nach Zustellzone. Vergleichswert aus dem Rechenkern in Spalte H."
wsn["A2"].font = F
head = ["Zone", "Felder", "Pakete", "Diesel-Vans (Standort)", "Fixkosten", "Variable Kosten inkl. Zweitzustellung", "Kosten pro Paket", "Rechenkern Runde 1", "Kapazitätsbedarf (Einheiten)"]
wsn.append([])
wsn.append(head)
style_header(wsn, 4, len(head))
zones = [("Stadt", "stadt", 7, 4), ("Agglo", "agglo", 11, 4), ("Land", "land", 15, 3)]
engine_r1 = {"stadt": 3.63, "agglo": 5.34, "land": 11.03}
for i, (label, z, fields, vans) in enumerate(zones):
    r = 5 + i
    wsn.cell(row=r, column=1, value=label)
    wsn.cell(row=r, column=2, value=fields).font = FINPUT
    wsn.cell(row=r, column=3, value=f"=B{r}*{P(f'zones.{z}.demand')}")
    wsn.cell(row=r, column=4, value=vans).font = FINPUT
    wsn.cell(row=r, column=5, value=f"=D{r}*({P('elements.vanDiesel.leasing')}+{P('elements.vanDiesel.driverCost')})")
    wsn.cell(row=r, column=6, value=f"=C{r}*{P('elements.vanDiesel.varCost')}*{P(f'zones.{z}.distance')}*(1+{P(f'zones.{z}.failRate')})")
    wsn.cell(row=r, column=7, value=f"=(E{r}+F{r})/C{r}")
    wsn.cell(row=r, column=8, value=engine_r1[z]).font = Font(name=FONT, size=10, color="808080")
    wsn.cell(row=r, column=9, value=f"=C{r}*{P(f'elements.vanDiesel.weight.{z}')}*(1+{P(f'zones.{z}.failRate')})")
r = 8
wsn.cell(row=r, column=1, value="Total").font = FB
for col in "BCDEF":
    wsn[f"{col}{r}"] = f"=SUM({col}5:{col}7)"
wsn[f"G{r}"] = f"=(E{r}+F{r})/C{r}"
wsn[f"I{r}"] = "=SUM(I5:I7)"
wsn["A10"] = "Kapazität der Diesel-Vans im Startnetz (Einheiten)"
wsn["G10"] = f"=D8*{P('elements.vanDiesel.capacity')}"
wsn["A11"] = "Auslastung des Startnetzes"
wsn["G11"] = "=I8/G10"
wsn["A12"] = "Anteil letzte Meile an den gesamten Paketkosten"
wsn["G12"] = f"=G8/(G8+{P('upstreamCostPerParcel')})"
wsn["A13"] = "Last-Mile-Tarif"
wsn["G13"] = f"={P('tariffPerParcel')}"
wsn["A14"] = "Abstand Startnetz zum Tarif (negativ = teurer als Tarif)"
wsn["G14"] = "=G13-G8"
wsn["A16"] = "Hinweis: Die Stadt-Vans in S1 bedienen auch die Agglo; die Zuordnung der Fixkosten folgt dem Standort (Spalte D). Kleine Abweichungen zu Spalte H entstehen, weil der Rechenkern die Pakete feldgenau zuteilt."
for row in wsn.iter_rows(min_row=5, max_row=16):
    for c in row:
        if c.font != FINPUT and c.font.color is None:
            c.font = F
for rr in range(5, 9):
    for cc in (5, 6):
        wsn.cell(row=rr, column=cc).number_format = "#,##0"
    wsn.cell(row=rr, column=7).number_format = "0.00"
    wsn.cell(row=rr, column=8).number_format = "0.00"
    wsn.cell(row=rr, column=9).number_format = "#,##0"
for cell in ("G10",):
    wsn[cell].number_format = "#,##0"
wsn["G11"].number_format = "0%"
wsn["G12"].number_format = "0%"
wsn["G13"].number_format = "0.00"
wsn["G14"].number_format = "0.00"
for cell in ("G13",):
    wsn[cell].font = FLINK
wsn["A16"].alignment = WRAP
set_widths(wsn, [52, 9, 10, 14, 12, 18, 14, 14, 16])

# ---------- Teststrategien ----------
wt = wb.create_sheet("Teststrategien")
wt["A1"] = f"Ergebnisse der Simulation: {SIM['nSeeds']} Ereignisfolgen pro Strategie, je ein Team allein, {CFG['rounds']} Runden"
wt["A1"].font = FB
cols = ["Schlüssel", "Strategie", "Kosten pro Paket (CHF)", "Pünktlichkeit", "CO₂ pro Paket (g)", "Fehlzustellungen", "Alle Standards erfüllt", "Standard landesweit", "Standard Kosten", "Standard alle Zonen", "Pünktl. Stadt", "Pünktl. Agglo", "Pünktl. Land", "Anteil letzte Meile", "Bewertung"]
wt.append([])
wt.append(cols)
style_header(wt, 3, len(cols))
strat_row = {}
for key, v in SIM["solo"].items():
    wt.append([key, v["label"], v["cost"], v["onTime"], v["co2"], v["fail"], v["allStd"], v["stdNat"], v["stdCost"], v["stdZones"], v["stadt"], v["agglo"], v["land"], v["share"]])
    r = wt.max_row
    strat_row[key] = r
    wt.cell(row=r, column=15, value=f'=IF(G{r}>=0.5,"trägt",IF(G{r}>0,"selten","scheitert"))')
    for c in range(1, 16):
        cell = wt.cell(row=r, column=c)
        cell.font = Font(name=FONT, size=10, color="595959") if 3 <= c <= 14 else F
        cell.border = BORDER
    wt.cell(row=r, column=3).number_format = "0.00"
    wt.cell(row=r, column=5).number_format = "0"
    for c in [4, 6, 7, 8, 9, 10, 11, 12, 13, 14]:
        wt.cell(row=r, column=c).number_format = "0%"
r0 = wt.max_row + 2
wt.cell(row=r0, column=1, value="Viererpartien (je Seed eine Partie)").font = FB
r = r0 + 1
for m in SIM["matches"]:
    wt.cell(row=r, column=1, value=" vs. ".join(m["keys"])).font = F
    wt.cell(row=r, column=2, value="Siege: " + ", ".join(f"{k} {n}" for k, n in m["wins"].items()) + f" · ohne Sieger: {m['noWinner']}").font = F
    r += 1
wt.cell(row=r + 1, column=1, value="Arbeitsmarkt: Kosten pro Paket eines Teams «Mix nach Zonen» mit und ohne ein Team, das massiv Vans einstellt").font = FB
wt.cell(row=r + 2, column=1, value="ohne").font = F
wt.cell(row=r + 2, column=3, value=SIM["labour"]["otherTeamCostWithout"]).number_format = "0.00"
wt.cell(row=r + 3, column=1, value="mit").font = F
wt.cell(row=r + 3, column=3, value=SIM["labour"]["otherTeamCostWith"]).number_format = "0.00"
wt.cell(row=r + 4, column=1, value="Mehrkosten für das andere Team").font = F
wt.cell(row=r + 4, column=3, value=f"=C{r + 3}/C{r + 2}-1").number_format = "0.0%"
wt.freeze_panes = "C4"
set_widths(wt, [16, 46, 12, 12, 11, 12, 12, 12, 11, 12, 10, 10, 10, 11, 11])

# ---------- Kalibrierung ----------
wk = wb.create_sheet("Kalibrierung")
wk["A1"] = "Kalibrierungskriterien aus dem Prüfdossier (Version 2.1) und ihre Prüfung"
wk["A1"].font = FH
wk.append([])
wk.append(["Nr.", "Kriterium", "Messung", "Wert", "Ergebnis"])
style_header(wk, 3, 5)
T = lambda key, col: f"Teststrategien!{col}{strat_row[key]}"  # noqa: E731
crit = [
    (1, "Startnetz verfehlt unter Belastung mindestens einen Standard", "Anteil Ereignisfolgen, in denen «Nichts ändern» alle Standards erfüllt", f"={T('nichts', 'G')}", '=IF(D{r}=0,"erfüllt","verfehlt")'),
    (2, "Mindestens zwei unterschiedliche Strategien erfüllen alle Standards", "Anzahl Strategien mit Erfolgsquote ≥ 50 %", f"=COUNTIF(Teststrategien!G4:G{wt.max_row},\">=0.5\")", '=IF(D{r}>=2,"erfüllt","verfehlt")'),
    (3, "Kein Element ist in allen Zonen das beste", "Hub nur Stadt, Box auf dem Land teuer, Post nur Land (Regel)", "Regel im Modell", "erfüllt"),
    (4, "Tarif wirtschaftlich begründet", "Paketpreis = Tarif + Vorlaufkosten", f"={P('tariffPerParcel')}+{P('upstreamCostPerParcel')}", '="Paketpreis CHF "&TEXT(D{r},"0.00")'),
    (5, "Anteil letzte Meile im Startnetz 40–50 %", "aus Blatt Startnetz", "=Startnetz!G12", '=IF(AND(D{r}>=0.4,D{r}<=0.5),"erfüllt","verfehlt")'),
    (6, "Rangfolge robust bei ±20 % je Parameter", "Variationen mit unveränderten Top 2 (von 16)", f"=COUNTIF(Sensitivität!F4:F{3 + len(SENS['rows'])},\"stabil\")", '=IF(D{r}>=14,"erfüllt","prüfen")'),
]
for i, (n, k, m, val, res) in enumerate(crit):
    r = 4 + i
    wk.cell(row=r, column=1, value=n)
    wk.cell(row=r, column=2, value=k)
    wk.cell(row=r, column=3, value=m)
    wk.cell(row=r, column=4, value=val)
    wk.cell(row=r, column=5, value=res.format(r=r) if "{r}" in res else res)
    for c in range(1, 6):
        cell = wk.cell(row=r, column=c)
        cell.font = FLINK if c == 4 and str(val).startswith("=") else F
        cell.alignment = WRAP
        cell.border = BORDER
wk["D4"].number_format = "0%"
wk["D8"].number_format = "0%"
wk["A11"] = "Befund Sensitivität: Steigen die Löhne um 20 %, erfüllt kaum noch eine Strategie den Kostenstandard. Wer Löhne erhöht, sollte den Tarif mitprüfen."
wk["A11"].font = F
set_widths(wk, [5, 46, 52, 14, 22])

# ---------- Sensitivität ----------
wsn2 = wb.create_sheet("Sensitivität")
wsn2["A1"] = f"Sensitivität: je ein Parameter um −20 % bzw. +20 % verändert, {SENS['seeds']} Ereignisfolgen pro Strategie"
wsn2["A1"].font = FB
wsn2.append([])
wsn2.append(["Parameter", "Faktor", "Beste Strategie", "Zweitbeste", "Startnetz scheitert weiter", "Top 2 unverändert"])
style_header(wsn2, 3, 6)
labels = SENS["labels"]
for row in SENS["rows"]:
    wsn2.append([row["label"], row["factor"], labels[row["best"]], labels[row["second"]], "ja" if row["startFails"] else "nein", "stabil" if row["top2Same"] else "Wechsel"])
    r = wsn2.max_row
    for c in range(1, 7):
        wsn2.cell(row=r, column=c).font = F
        wsn2.cell(row=r, column=c).border = BORDER
    wsn2.cell(row=r, column=2).number_format = "0%"
r = wsn2.max_row + 2
wsn2.cell(row=r, column=1, value="Basis (unverändert): " + " > ".join(labels[k] for k in SENS["base"]["ranking"][:3])).font = F
set_widths(wsn2, [30, 9, 44, 44, 20, 16])

out = ROOT / "docs/Balancing.xlsx"
wb.save(out)
print(out)
