"""Übernimmt die Werte aus dem Blatt «Parameter» der Balancing-Excel in config/params.json.

Aufruf: python3 tools/excel_to_params.py docs/Balancing.xlsx
Danach: npm test && npm run simulate (prüfen), dann veröffentlichen.
"""
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
xlsx = Path(sys.argv[1] if len(sys.argv) > 1 else ROOT / "docs/Balancing.xlsx")
cfg_path = ROOT / "config/params.json"
cfg = json.loads(cfg_path.read_text(encoding="utf-8"))

wb = load_workbook(xlsx, data_only=True)
ws = wb["Parameter"]
changed = []
for row in ws.iter_rows(min_row=2, values_only=True):
    key, value = row[0], row[3]
    if not key or value is None:
        continue
    parts = str(key).split(".")
    node = cfg
    for p in parts[:-1]:
        node = node[p]
    old = node[parts[-1]]
    new = type(old)(value) if not isinstance(old, str) else str(value)
    if isinstance(old, float) or isinstance(new, float):
        new = float(value)
    if new != old:
        node[parts[-1]] = new
        changed.append((key, old, new))

cfg_path.write_text(json.dumps(cfg, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
for k, o, n in changed:
    print(f"{k}: {o} -> {n}")
print(f"{len(changed)} Werte geändert.")
